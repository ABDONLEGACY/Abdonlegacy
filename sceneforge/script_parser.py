"""Read a script (.docx / .txt / .md or pasted text) and split it into scenes."""
import re
from pathlib import Path

# "Scene 1:", "SCENE 2 -", "**Scene 3.**", "## Scene 4", "Scene 5)" ...
SCENE_RE = re.compile(
    r"^[\s#>*_\-]*scene\s*(\d+)\s*(?:[:.\-–—)\]]+|\*\*|\s|$)\s*[*_]*\s*(.*)$",
    re.IGNORECASE,
)
# Things in [square brackets] or {curly braces} are treated as notes, not narration.
NOTE_RE = re.compile(r"\[[^\]]*\]|\{[^}]*\}")
MD_RE = re.compile(r"(\*\*|__|`|^#+\s*|^>\s*)", re.MULTILINE)


def read_script_file(path) -> str:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".docx":
        import docx
        doc = docx.Document(str(path))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    lines.append(cell.text)
        return "\n".join(lines)
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def clean_narration(text: str) -> str:
    text = NOTE_RE.sub(" ", text)
    text = MD_RE.sub("", text)
    text = text.replace(" ", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def parse_scenes(text: str):
    """Return (scenes, warnings). Each scene = {"number": int, "text": str}."""
    warnings = []
    scenes = []
    current = None
    preamble = []
    for line in text.splitlines():
        m = SCENE_RE.match(line)
        if m:
            if current is not None:
                scenes.append(current)
            current = {"number": int(m.group(1)), "lines": [m.group(2)]}
        elif current is not None:
            current["lines"].append(line)
        elif line.strip():
            preamble.append(line.strip())
    if current is not None:
        scenes.append(current)

    out = []
    for s in scenes:
        narration = clean_narration("\n".join(s["lines"]))
        out.append({"number": s["number"], "text": narration})

    if not out:
        warnings.append('No scenes found. Each scene must start with a label like "Scene 1:" on its own line.')
        return out, warnings
    if preamble:
        warnings.append(f"Text before 'Scene {out[0]['number']}' was ignored: \"{' '.join(preamble)[:80]}...\"")
    nums = [s["number"] for s in out]
    seen, dups = set(), []
    for n in nums:
        if n in seen:
            dups.append(n)
        seen.add(n)
    if dups:
        warnings.append(f"Scene numbers used more than once: {sorted(set(dups))}")
    expected = list(range(nums[0], nums[0] + len(nums)))
    if nums != expected and not dups:
        missing = sorted(set(range(1, max(nums) + 1)) - set(nums))
        if missing:
            warnings.append(f"Scene numbers missing from the script: {missing[:20]}")
        if nums != sorted(nums):
            warnings.append("Scene numbers are not in order in the script. They are used in the order written.")
    empty = [s["number"] for s in out if not s["text"]]
    if empty:
        warnings.append(f"These scenes have no narration text: {empty}")
    return out, warnings


def natural_key(s: str):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", str(s))]
